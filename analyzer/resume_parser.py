import os
import re
from PyPDF2 import PdfReader
from docx import Document


def clean_text(text):
    if not text:
        return ""
    # Replace non-breaking spaces and irregular whitespace
    text = text.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    # Remove control characters except standard line breaks and tabs
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    return text


def extract_text(file_path):
    if not os.path.exists(file_path):
        return ""

    extension = os.path.splitext(file_path)[1].lower()

    try:
        if extension == ".pdf":
            return clean_text(extract_pdf_text(file_path))
        elif extension == ".docx":
            return clean_text(extract_docx_text(file_path))
    except Exception as e:
        print(f"Error parsing {file_path}: {e}")
        return ""

    return ""


def extract_pdf_text(file_path):
    text = []
    try:
        reader = PdfReader(file_path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                return ""

        for page in reader.pages:
            page_text = page.extract_text()
            if page_text and page_text.strip():
                text.append(page_text.strip())

        return "\n\n".join(text)
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
        return ""


def extract_docx_text(file_path):
    text = []
    try:
        document = Document(file_path)

        for paragraph in document.paragraphs:
            if paragraph.text and paragraph.text.strip():
                text.append(paragraph.text.strip())

        # Also extract text inside tables (resumes often use tables for layout)
        for table in document.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text.append(" | ".join(row_texts))

        return "\n".join(text)
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
        return ""
